from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient

from bidmaster.api.tenderbrief import _lot_aliases, _lot_matches, _lot_options, create_router
from bidmaster.orchestration.pipeline import Pipeline
from bidmaster.schemas.document import ParsedDocument, TableModel
from bidmaster.schemas.evidence import EvidenceStore
from bidmaster.schemas.fields import FieldExtraction
from bidmaster.schemas.report import TenderReport
from bidmaster.summary.workbook import build_summary_xlsx
from bidmaster.structure.lot_scope import extract_lot_table_fields


def test_summary_tables_preserve_package_rows_and_split_qualification_scope():
    parsed = ParsedDocument(
        doc_id="a1b2c3d4e5f6",
        file_name="招标资料.docx",
        file_type="docx",
        tables=[
            TableModel(table_id="split", page_nos=[2], rows=[
                ["分标编号", "分标名称", "分包数量", "是否允许分包"],
                ["062639-001", "工程预结算审核", "2", "不允许"],
            ]),
            TableModel(table_id="packages", page_nos=[3], rows=[
                ["分标名称", "包号", "包名称", "项目名称", "网省采购申请号", "招标范围",
                 "建设地点", "工期/服务期", "单项最高限价（万元）", "包最高限价（万元）"],
                ["工程预结算审核", "包01", "济南项目审核", "济南工程", "3901001",
                 "见技术规范书", "山东省济南市", "自合同签订起至竣工结算审核完成",
                 "299.27", "299.27"],
                ["工程预结算审核", "包02", "泰安项目审核", "泰安工程", "3901002",
                 "见技术规范书", "山东省泰安市", "自合同签订起至竣工结算审核完成",
                 "165.32", "165.32"],
            ]),
            TableModel(table_id="qualifications", page_nos=[5], rows=[
                ["分标名称", "包号", "资质要求", "业绩要求", "人员要求"],
                ["工程预结算审核", "包01-包02", "/", "近三年同类业绩", "一级注册造价工程师"],
            ]),
        ],
    )
    fields = extract_lot_table_fields(parsed, EvidenceStore(doc_id=parsed.doc_id))

    assert set(fields) >= {"包01", "包02"}
    first = fields["包01"]
    second = fields["包02"]
    assert first["split_no"].value_raw == "062639-001"
    assert first["package_count"].value_raw == "2"
    assert first["subcontract_allowed"].value_raw == "不允许"
    assert first["package_name"].value_raw == "济南项目审核"
    assert first["project_name"].value_raw == "济南工程"
    assert first["application_no"].value_raw == "3901001"
    assert first["single_price_limit"].value_raw == "299.27万元"
    assert first["package_price_limit"].value_raw == "299.27万元"
    assert first["special_performance"].value_raw == "近三年同类业绩"
    assert second["package_name"].value_raw == "泰安项目审核"
    assert second["package_price_limit"].value_raw == "165.32万元"
    assert second["special_personnel"].value_raw == "一级注册造价工程师"

    options = _lot_options(TenderReport(doc_id=parsed.doc_id, lot_fields=fields))
    assert [item["id"] for item in options] == ["包01", "包02"]
    package_one = _lot_aliases("包01", options[0]["name"])
    assert _lot_matches("济南项目审核包1", package_one)
    assert not _lot_matches("泰安项目审核包2", package_one)


def test_summary_workbook_exports_one_package_and_directory():
    report = TenderReport(
        doc_id="a1b2c3d4e5f6", file_name="测试招标文件.zip",
        fields={"project_name": FieldExtraction(
            field_key="project_name", field_label="项目名称", status="found",
            value_raw="国网济南工程", value_normalized="国网济南工程")},
        lot_fields={"包01": {
            "lot_no": FieldExtraction(field_key="lot_no", field_label="包号",
                                       status="found", value_raw="包01", value_normalized="包01"),
            "package_name": FieldExtraction(field_key="package_name", field_label="包名称",
                                              status="found", value_raw="济南项目审核",
                                              value_normalized="济南项目审核"),
            "split_name": FieldExtraction(field_key="split_name", field_label="分标名称",
                                            status="found", value_raw="工程预结算审核",
                                            value_normalized="工程预结算审核"),
        }},
    )
    content = build_summary_xlsx(report, "包01", {
        "business": [{"title": "商务偏差表", "evidence_id": "ev-1", "origin": "投标文件格式"}],
        "technical": [{"title": "施工方案", "evidence_id": "ev-2", "origin": "评分项补充",
                       "status": "建议加入目录，需核对"}],
        "price": [],
    })
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    assert workbook.sheetnames == ["投标要点汇总表", "投标文件目录", "评分项"]
    summary_values = [cell for row in workbook["投标要点汇总表"].iter_rows(values_only=True)
                      for cell in row if isinstance(cell, str)]
    directory_values = [cell for row in workbook["投标文件目录"].iter_rows(values_only=True)
                        for cell in row if isinstance(cell, str)]
    assert "包01" in summary_values
    assert "济南项目审核" in summary_values
    assert "商务偏差表" in directory_values
    assert "施工方案" in directory_values
    workbook.close()


def test_export_route_returns_download_for_selected_package(tmp_path: Path):
    pipeline = Pipeline(tmp_path / "work")
    job_id = "a1b2c3d4e5f6"
    job_dir = pipeline.work_root / job_id
    job_dir.mkdir(parents=True)
    report = TenderReport(
        doc_id=job_id, file_name="项目资料.zip",
        lot_fields={"包01": {
            "lot_no": FieldExtraction(field_key="lot_no", field_label="包号",
                                       status="found", value_raw="包01", value_normalized="包01"),
            "package_name": FieldExtraction(field_key="package_name", field_label="包名称",
                                              status="found", value_raw="济南审核包",
                                              value_normalized="济南审核包"),
        }},
    )
    (job_dir / "06_report.json").write_text(
        json.dumps(report.model_dump(mode="json"), ensure_ascii=False), encoding="utf-8")
    app = FastAPI()
    app.include_router(create_router(pipeline))

    with TestClient(app) as client:
        response = client.get(f"/api/jobs/{job_id}/summary.xlsx?lot_id=%E5%8C%8501")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    assert workbook.sheetnames[0] == "投标要点汇总表"
    workbook.close()
